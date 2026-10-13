"""Exports of a run's approved article: Markdown, HTML, DOCX and the hero as the CMS WebP (M7)."""

import io
from dataclasses import dataclass
from html.parser import HTMLParser

import yaml
from django.utils.html import escape
from django.utils.text import slugify

from . import markup

LANG = {"US": "en-US", "UK": "en-GB", "AU": "en-AU", "CA": "en-CA", "NZ": "en-NZ", "EU": "en"}
HERO_SIZE = (720, 405)  # what the CMS serves (sds-hero-stock)


class _ExportResolver(markup.WagtailResolver):
    """Absolute URLs (exports leave the admin), and remembers which image each URL came from."""

    def __init__(self, **kwargs):
        from django.conf import settings

        super().__init__(base_url=settings.WAGTAILADMIN_BASE_URL, **kwargs)
        self.images_by_url = {}

    def images(self, ids):
        from wagtail.images import get_image_model

        found = super().images(ids)
        library = get_image_model().objects.in_bulk(found.keys())
        self.images_by_url.update({resolved.url: library[pk] for pk, resolved in found.items()})
        return found


def _resolver():
    return _ExportResolver()


class ArticleError(Exception):
    """The run has nothing approved to export or publish yet; the message says what's missing."""


@dataclass
class ArticleContent:
    text: str
    hero: object  # a Wagtail image or None
    region: str
    variant: str  # "final" | "localised"


def _approved_version(run, key):
    step = run.steps.filter(key=key, status="approved").first()
    return step.latest_version if step else None


def article_content(run, variant="final"):
    if variant == "localised":
        version = _approved_version(run, "localise")
        region = run.brief.localise_to
        if not version or not version.content.strip():
            raise ArticleError("The localised variant isn't approved yet.")
    else:
        version = next(
            (v for key in ("cleanup", "review_edit", "draft") if (v := _approved_version(run, key))),
            None,
        )
        region = run.brief.region
        if not version or not version.content.strip():
            raise ArticleError("There is no approved article yet. Approve the draft first.")
    hero_version = _approved_version(run, "hero_image")
    return ArticleContent(
        text=version.content,
        hero=hero_version.image if hero_version else None,
        region=region,
        variant=variant,
    )


def _selected(run, kind):
    option = run.metadata_options.filter(kind=kind, selected=True).first()
    return " ".join(option.text.split()).strip("`") if option else ""


def article_meta(run, content):
    """Page title (the H1), SEO title, meta description and slug from the run's selections."""
    h1, _ = markup.split_h1(content.text)
    seo_title = _selected(run, "title")
    title = h1 or seo_title or run.brief.topic
    return {
        "title": title,
        "seo_title": seo_title or title,
        "meta_description": _selected(run, "meta"),
        "slug": slugify(_selected(run, "slug") or seo_title or title)[:80] or "article",
    }


def filename(run, variant="final"):
    content = article_content(run, variant)
    slug = article_meta(run, content)["slug"]
    return f"{slug}-{content.region.lower()}" if variant == "localised" else slug


def markdown_export(run, variant="final"):
    content = article_content(run, variant)
    meta = article_meta(run, content)
    front = {
        **meta,
        "region": content.region,
        "primary_keyword": run.brief.pkeyword,
        "secondary_keywords": list(run.brief.skeywords or []),
        "hero_image": f"{meta['slug']}.webp" if content.hero else None,
        "studio_run": run.pk,
    }
    front_matter = yaml.safe_dump(front, sort_keys=False, allow_unicode=True).strip()
    body = markup.resolve_references(markup.strip_workflow_sections(content.text), _resolver())
    return f"---\n{front_matter}\n---\n\n{body}\n"


def html_export(run, variant="final"):
    content = article_content(run, variant)
    meta = article_meta(run, content)
    body = markup.markdown_to_html(markup.strip_workflow_sections(content.text), _resolver())
    hero = (
        f'<img src="{escape(meta["slug"])}.webp" alt="" width="720" height="405">\n' if content.hero else ""
    )
    return (
        "<!doctype html>\n"
        f'<html lang="{LANG.get(content.region, "en")}">\n<head>\n<meta charset="utf-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"<title>{escape(meta['seo_title'])}</title>\n"
        f'<meta name="description" content="{escape(meta["meta_description"])}">\n'
        f"</head>\n<body>\n<article>\n{hero}{body}\n</article>\n</body>\n</html>\n"
    )


SVG_RASTER_WIDTH = 1440  # DOCX and Pillow can't use SVG; it's drawn at this width first


def _open_picture(image, handle):
    from PIL import Image

    if not image.is_svg():
        return Image.open(handle)
    import cairosvg

    # unsafe=False (the default): no external files or URLs are fetched while drawing.
    png = cairosvg.svg2png(file_obj=handle, output_width=SVG_RASTER_WIDTH, background_color="white")
    return Image.open(io.BytesIO(png))


def _image_bytes(image, size=None, fmt="PNG"):
    from PIL import Image, ImageOps

    with image.file.open("rb") as handle, _open_picture(image, handle) as source:
        picture = source.convert("RGB")
        if size:
            picture = ImageOps.fit(picture, size, Image.LANCZOS)
        buffer = io.BytesIO()
        picture.save(buffer, fmt, **({"quality": 90} if fmt == "WEBP" else {}))
    return buffer.getvalue()


def hero_webp(run):
    """The hero at 720x405 WebP, as the sds-hero-stock skill delivers it for the CMS."""
    hero = article_content(run).hero
    if hero is None:
        raise ArticleError("There is no approved hero image.")
    return _image_bytes(hero, HERO_SIZE, "WEBP")


# ---- DOCX ----


def _add_hyperlink(paragraph, url, text):
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    text_node.set(qn("xml:space"), "preserve")
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


class _DocxWriter(HTMLParser):
    """Walks the sanitised article HTML and writes Word paragraphs, lists, tables and links."""

    def __init__(self, document, images_by_url=None):
        super().__init__(convert_charrefs=True)
        self.doc = document
        self.images_by_url = images_by_url or {}
        self.paragraph = None
        self.bold = self.italic = self.code = self.quote = self.underline = self.strike = 0
        self.caption = False
        self.link = None
        self.lists = []
        self.rows = self.row = self.cell = None

    def _picture(self, src):
        from docx.shared import Inches

        image = self.images_by_url.get(src)
        if image is not None:
            self.doc.add_picture(io.BytesIO(_image_bytes(image)), width=Inches(6.5))

    def _new_paragraph(self, style=None):
        self.paragraph = self.doc.add_paragraph(style=style)

    def handle_starttag(self, tag, attrs):
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.paragraph = self.doc.add_heading(level=min(int(tag[1]), 4))
        elif tag == "p":
            if not (self.lists and self.paragraph is not None and not self.paragraph.text):
                self._new_paragraph("Quote" if self.quote else None)
        elif tag in {"ul", "ol"}:
            self.lists.append(tag)
        elif tag == "li":
            self._new_paragraph("List Number" if self.lists and self.lists[-1] == "ol" else "List Bullet")
        elif tag in {"strong", "b"}:
            self.bold += 1
        elif tag in {"em", "i"}:
            self.italic += 1
        elif tag == "code":
            self.code += 1
        elif tag == "u":
            self.underline += 1
        elif tag == "s":
            self.strike += 1
        elif tag == "img":
            self.paragraph = None
            self._picture(dict(attrs).get("src", ""))
        elif tag == "figcaption":
            self._new_paragraph("Caption")
            self.caption = True
        elif tag == "blockquote":
            self.quote += 1
        elif tag == "a":
            self.link = dict(attrs).get("href")
        elif tag == "br" and self.paragraph is not None:
            self.paragraph.add_run().add_break()
        elif tag == "table":
            self.rows = []
        elif tag == "tr" and self.rows is not None:
            self.row = []
        elif tag in {"th", "td"} and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "figcaption"}:
            self.paragraph = None
            self.caption = False if tag == "figcaption" else self.caption
        elif tag == "u":
            self.underline -= 1
        elif tag == "s":
            self.strike -= 1
        elif tag in {"ul", "ol"} and self.lists:
            self.lists.pop()
        elif tag in {"strong", "b"}:
            self.bold -= 1
        elif tag in {"em", "i"}:
            self.italic -= 1
        elif tag == "code":
            self.code -= 1
        elif tag == "blockquote":
            self.quote -= 1
        elif tag == "a":
            self.link = None
        elif tag in {"th", "td"} and self.cell is not None:
            self.row.append(("".join(self.cell).strip(), tag == "th"))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag == "table" and self.rows is not None:
            self._write_table(self.rows)
            self.rows = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)
            return
        if self.paragraph is None:
            if not data.strip():
                return
            self._new_paragraph()
        if self.link:
            _add_hyperlink(self.paragraph, self.link, data)
            return
        run = self.paragraph.add_run(data)
        run.bold = bool(self.bold) or None
        run.italic = bool(self.italic) or None
        run.underline = bool(self.underline) or None
        run.font.strike = bool(self.strike) or None
        if self.code:
            run.font.name = "Consolas"

    def _write_table(self, rows):
        if not rows:
            return
        width = max(len(row) for row in rows)
        table = self.doc.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, (text, header) in enumerate(row):
                cell = table.cell(r, c)
                cell.text = text
                if header:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True


def docx_export(run, variant="final"):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    content = article_content(run, variant)
    meta = article_meta(run, content)
    document = Document()
    document.core_properties.title = meta["seo_title"]
    document.core_properties.subject = meta["meta_description"]
    document.core_properties.keywords = run.brief.pkeyword

    details = document.add_paragraph()
    for label, value in (
        ("SEO title", meta["seo_title"]),
        ("Meta description", meta["meta_description"] or "(none selected)"),
        ("Slug", meta["slug"]),
        ("Region", content.region),
    ):
        line = details.add_run(f"{label}: {value}\n")
        line.font.size = Pt(9)
        line.font.color.rgb = RGBColor(0x62, 0x62, 0x62)
    if content.hero is not None:
        document.add_picture(io.BytesIO(_image_bytes(content.hero)), width=Inches(6.5))

    resolver = _resolver()
    body = markup.markdown_to_html(markup.strip_workflow_sections(content.text), resolver)
    writer = _DocxWriter(document, resolver.images_by_url)
    writer.feed(body)
    writer.close()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
