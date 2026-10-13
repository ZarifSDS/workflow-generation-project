"""M7: ArticlePage drafts, moderation, exports and the publish webhook."""

import hashlib
import hmac
import io
import json

import httpx
import pytest
import respx
import yaml
from django.core.files.images import ImageFile
from django.urls import reverse
from PIL import Image
from wagtail.images import get_image_model

from apps.content import exports, markup, publishing
from apps.content.models import ArticleIndexPage, ArticlePage
from apps.studio import services
from apps.studio.models import ContentBrief, MetadataOption, StepRun
from apps.studio.tasks import execute_step

pytestmark = pytest.mark.django_db

ARTICLE = """# SDS Retention Requirements: How Long to Keep SDSs

OSHA treats safety data sheets as **exposure records**. Keep them for *30 years*.
See [OSHA 1910.1020](https://www.osha.gov/laws-regs/regulations/standardnumber/1910/1910.1020).

## What the rule says

- Keep the current SDS on every shift.
- Archive old versions.

### Who must comply

Employers with hazardous chemicals.

| Record | Keep for |
|---|---|
| SDS | 30 years |
| Exposure log | 30 years |

## Changelog
- Removed filler.
"""


def _approve(run, key, content="", image=None):
    step = run.steps.get(key=key)
    services.add_version(step, content, image=image)
    step.status = StepRun.Status.APPROVED
    step.save()
    return step


def _hero_image():
    buffer = io.BytesIO()
    Image.new("RGB", (1440, 810), "white").save(buffer, "PNG")
    return get_image_model().objects.create(title="Hero", file=ImageFile(buffer, name="hero.png"))


@pytest.fixture
def run(admin_user):
    brief = ContentBrief.objects.create(
        topic="How long to keep SDSs",
        pkeyword="SDS retention requirements",
        skeywords=["how long to keep safety data sheets"],
        region="US",
        localise_to="UK",
        content_type="seo_blog",
        created_by=admin_user,
    )
    run = services.create_run(brief, {"localise": {"enabled": True}}, user=admin_user)
    _approve(run, "draft", ARTICLE)
    _approve(run, "cleanup", ARTICLE)
    _approve(run, "hero_image", "Branded hero.", image=_hero_image())
    for kind, text in (
        ("title", "SDS Retention Requirements: How Long to Keep SDSs"),
        ("meta", "OSHA treats SDSs as exposure records. Learn the retention rules."),
        ("slug", "`sds-retention-requirements`"),
    ):
        MetadataOption.objects.create(run=run, kind=kind, text=text, selected=True)
    return run


def _run_publish(run):
    step = run.steps.get(key="publish")
    step.status = StepRun.Status.QUEUED
    step.save()
    execute_step(step.pk)
    step.refresh_from_db()
    return step


# ---- Markdown conversion ----


def test_to_stream_maps_headings_tables_and_rich_text():
    _, body = markup.split_h1(ARTICLE)
    blocks = markup.to_stream(body)
    types = [block["type"] for block in blocks]
    assert types == ["paragraph", "heading", "paragraph", "heading", "paragraph", "table"]
    assert blocks[1]["value"] == {"text": "What the rule says", "level": "h2", "anchor": "what-the-rule-says"}
    assert blocks[3]["value"]["level"] == "h3"
    assert "<b>exposure records</b>" in blocks[0]["value"]
    assert "<i>30 years</i>" in blocks[0]["value"]
    assert 'href="https://www.osha.gov' in blocks[0]["value"]
    assert "<ul>" in blocks[2]["value"]
    assert blocks[5]["value"]["data"][0] == ["Record", "Keep for"]
    assert blocks[5]["value"]["first_row_is_table_header"] is True
    assert "Changelog" not in json.dumps(blocks)


def test_untrusted_html_is_stripped():
    blocks = markup.to_stream("Hi <script>alert(1)</script> [x](javascript:alert(1)) <img src=x onerror=1>")
    assert "script" not in blocks[0]["value"]
    assert "javascript" not in blocks[0]["value"]
    assert "<img" not in blocks[0]["value"]


def test_split_h1_and_verify_markers():
    h1, body = markup.split_h1(ARTICLE)
    assert h1 == "SDS Retention Requirements: How Long to Keep SDSs"
    assert not body.startswith("#")
    assert markup.verify_markers(ARTICLE + "\nOne more claim [VERIFY].") == 1


# ---- Publish step ----


def test_publish_creates_draft_page_in_moderation(run):
    step = _run_publish(run)
    assert step.status == StepRun.Status.READY, step.error
    page = run.article_page
    assert page.live is False
    assert page.title == "SDS Retention Requirements: How Long to Keep SDSs"
    assert page.slug == "sds-retention-requirements"
    assert page.search_description.startswith("OSHA treats SDSs")
    assert page.hero_image is not None
    assert page.get_parent().specific_class is ArticleIndexPage
    assert page.workflow_in_progress
    assert page.current_workflow_state.workflow.name == "Moderators approval"
    assert [block.block_type for block in page.body][:1] == ["heading"]
    assert page.intro.startswith("OSHA treats safety data sheets as exposure records.")
    assert page.category == "article"
    draft = page.get_latest_revision_as_object()  # tags live in the draft revision until publishing
    assert set(draft.tags.names()) == {"SDS retention requirements", "how long to keep safety data sheets"}
    version = step.latest_version
    assert "Moderators approval" in version.content
    assert version.metrics["publish"]["page_id"] == page.pk


def test_publishing_again_updates_the_same_page(run):
    _run_publish(run)
    page = run.article_page
    services.request_regenerate(run.steps.get(key="publish"), "ollama::glm-4.7-flash:latest")
    step = _run_publish(run)
    assert step.status == StepRun.Status.READY, step.error
    assert ArticlePage.objects.count() == 1
    assert page.revisions.count() == 2
    assert "already in the" in step.latest_version.content


def test_slug_is_made_unique_among_articles(run, admin_user):
    index = publishing.article_index()
    index.add_child(instance=ArticlePage(title="Older", slug="sds-retention-requirements", body=[]))
    _run_publish(run)
    assert run.article_page.slug == "sds-retention-requirements-2"


def test_publish_without_an_approved_article_fails_clearly(admin_user):
    brief = ContentBrief.objects.create(topic="Empty", pkeyword="empty", created_by=admin_user)
    run = services.create_run(brief, {})
    step = _run_publish(run)
    assert step.status == StepRun.Status.FAILED
    assert "no approved article" in step.error


def test_published_page_renders(run, client):
    _run_publish(run)
    run.article_page.get_latest_revision().publish()
    response = client.get(run.article_page.url)
    assert response.status_code == 200
    html = response.text
    assert '<h2 id="what-the-rule-says">What the rule says</h2>' in html
    assert "OSHA treats safety data sheets" in html  # the intro
    assert "Exposure log" in html
    assert "exposure records" in html


@respx.mock
def test_webhook_gets_signed_payload(run, settings):
    settings.STUDIO_PUBLISH_WEBHOOK_URL = "http://n8n.test/webhook/studio"
    settings.STUDIO_WEBHOOK_SECRET = "s3cret"
    hook = respx.post("http://n8n.test/webhook/studio").mock(return_value=httpx.Response(200))
    step = _run_publish(run)
    request = hook.calls.last.request
    payload = json.loads(request.content)
    expected = hmac.new(b"s3cret", request.content, hashlib.sha256).hexdigest()
    assert request.headers["X-Studio-Signature"] == f"sha256={expected}"
    assert payload["event"] == "article.draft_saved"
    assert payload["slug"] == "sds-retention-requirements"
    assert payload["markdown"].startswith("---\n")
    assert payload["admin_url"].endswith(f"/admin/pages/{run.article_page.pk}/edit/")
    assert "Sent to the webhook" in step.latest_version.note


@respx.mock
def test_webhook_failure_does_not_fail_the_step(run, settings):
    settings.STUDIO_PUBLISH_WEBHOOK_URL = "http://n8n.test/webhook/studio"
    respx.post("http://n8n.test/webhook/studio").mock(side_effect=httpx.ConnectError("down"))
    step = _run_publish(run)
    assert step.status == StepRun.Status.READY
    assert "couldn't be reached" in step.latest_version.note


# ---- Exports ----


def test_markdown_export_has_front_matter(run):
    text = exports.markdown_export(run)
    _, front, body = text.split("---\n", 2)
    meta = yaml.safe_load(front)
    assert meta["slug"] == "sds-retention-requirements"
    assert meta["seo_title"] == "SDS Retention Requirements: How Long to Keep SDSs"
    assert meta["hero_image"] == "sds-retention-requirements.webp"
    assert body.strip().startswith("# SDS Retention Requirements")
    assert "Changelog" not in body


def test_html_export_is_a_standalone_document(run):
    html = exports.html_export(run)
    assert html.startswith("<!doctype html>")
    assert '<html lang="en-US">' in html
    assert '<meta name="description" content="OSHA treats SDSs' in html
    assert "<table>" in html


def test_docx_export_has_headings_lists_table_and_link(run):
    from docx import Document

    document = Document(io.BytesIO(exports.docx_export(run)))
    styles = [(p.style.name, p.text) for p in document.paragraphs]
    assert ("Heading 1", "SDS Retention Requirements: How Long to Keep SDSs") in styles
    assert ("Heading 2", "What the rule says") in styles
    assert ("List Bullet", "Keep the current SDS on every shift.") in styles
    assert document.tables[0].cell(1, 1).text == "30 years"
    assert len(document.inline_shapes) == 1  # the hero
    assert any("osha.gov" in rel.target_ref for rel in document.part.rels.values() if rel.is_external)


def test_hero_webp_is_cms_size(run):
    with Image.open(io.BytesIO(exports.hero_webp(run))) as image:
        assert image.format == "WEBP"
        assert image.size == (720, 405)


def test_export_views_download_and_explain_missing_variants(run, admin_client):
    response = admin_client.get(reverse("studio:run_export", args=[run.pk, "md"]))
    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="sds-retention-requirements.md"'
    response = admin_client.get(reverse("studio:run_export", args=[run.pk, "docx"]) + "?variant=localised")
    assert response.status_code == 302
    _approve(run, "localise", "# SDS Retention Requirements (UK)\n\nCOSHH text.")
    response = admin_client.get(reverse("studio:run_export", args=[run.pk, "docx"]) + "?variant=localised")
    assert response["Content-Disposition"] == 'attachment; filename="sds-retention-requirements-uk.docx"'
    assert admin_client.get(reverse("studio:run_export", args=[run.pk, "pdf"])).status_code == 404


def test_workspace_publish_step_links_page_and_exports(run, admin_client):
    _run_publish(run)
    response = admin_client.get(reverse("studio:run", args=[run.pk]) + "?step=publish")
    html = response.text
    assert reverse("wagtailadmin_pages:edit", args=[run.article_page.pk]) in html
    assert "Moderators approval" in html
    assert reverse("studio:run_export", args=[run.pk, "docx"]) in html
    assert "Hero WebP" in html
    assert "Primary keyword in title" not in html


def test_page_editor_links_back_to_the_run(run, admin_client):
    _run_publish(run)
    response = admin_client.get(reverse("wagtailadmin_pages:edit", args=[run.article_page.pk]))
    assert response.status_code == 200
    assert reverse("studio:run", args=[run.pk]) in response.text


def test_run_again_button_requeues_publish(run, admin_client):
    _run_publish(run)
    url = reverse("studio:run", args=[run.pk]) + "?step=publish"
    assert "Run publish again" in admin_client.get(url).text
    admin_client.post(url, {"step": "publish", "action": "rerun"})
    assert run.steps.get(key="publish").status == StepRun.Status.QUEUED
    with pytest.raises(ValueError):
        services.rerun_step(run.steps.get(key="draft"))


# ---- M9 phase 2: figures, embeds, FAQ, intro, author ----

RICH_ARTICLE = """# Management of change

A change can bring hazards the original risk assessment missed.

## What counts {#what-counts}

Some ++temporary bypasses++ and ~~old~~ new rules. See [the SDS page](wagtail://page?id={page_id}).

![Engineer at the panel](wagtail://image?id={image_id} "Every change needs a reviewer.")

![MOC explained](wagtail://media?url=https%3A%2F%2Fwww.youtube.com%2Fwatch%3Fv%3Dabc123)

:::faq
### Does a temporary change need MOC?

Yes, with an end date.
:::
"""


def _rich_run(run):
    image = _hero_image()
    page_id = publishing.article_index().pk
    text = RICH_ARTICLE.replace("{page_id}", str(page_id)).replace("{image_id}", str(image.pk))
    for key in ("draft", "cleanup"):
        step = run.steps.get(key=key)
        services.add_version(step, text)
    return image, page_id


def test_build_article_maps_every_block_type(run):
    image, page_id = _rich_run(run)
    _, body = markup.split_h1(run.steps.get(key="cleanup").latest_version.content)
    article = markup.build_article(body)
    assert article.intro == "A change can bring hazards the original risk assessment missed."
    types = [block["type"] for block in article.blocks]
    assert types == ["heading", "paragraph", "figure", "embed", "faq"]
    assert article.blocks[0]["value"]["anchor"] == "what-counts"
    paragraph = article.blocks[1]["value"]
    assert "<u>temporary bypasses</u>" in paragraph and "<s>old</s>" in paragraph
    assert (
        f'<a id="{page_id}" linktype="page">' in paragraph or f'linktype="page" id="{page_id}"' in paragraph
    )
    figure = article.blocks[2]["value"]
    assert figure["image"] == {"image": image.pk, "decorative": False, "alt_text": "Engineer at the panel"}
    assert figure["caption"] == "Every change needs a reviewer."
    assert article.blocks[3]["value"] == "https://www.youtube.com/watch?v=abc123"
    faq = article.blocks[4]["value"]["items"][0]
    assert faq["question"] == "Does a temporary change need MOC?" and "end date" in faq["answer"]
    assert article.warnings == []


def test_missing_image_is_a_publish_warning(run):
    article = markup.build_article('## Heading\n\n![Gone](wagtail://image?id=99999 "Caption")')
    assert [block["type"] for block in article.blocks] == ["heading"]
    assert any("99999" in warning for warning in article.warnings)


def test_published_page_shows_figure_faq_and_author(run, client, admin_user):
    from apps.content.models import Author

    _rich_run(run)
    Author.objects.create(name="Dana Okafor", role="EHS consultant", user=admin_user)
    _run_publish(run)
    page = run.article_page
    assert page.author.name == "Dana Okafor"
    page.get_latest_revision().publish()
    html = client.get(page.url).text
    assert 'alt="Engineer at the panel"' in html
    assert "<figcaption>Every change needs a reviewer.</figcaption>" in html
    assert "Does a temporary change need MOC?" in html
    assert "By Dana Okafor, EHS consultant" in html
    assert 'id="what-counts"' in html


def test_exports_resolve_references(run):
    from docx import Document

    image, page_id = _rich_run(run)
    markdown = exports.markdown_export(run)
    assert "wagtail://" not in markdown
    assert "[MOC explained](https://www.youtube.com/watch?v=abc123)" in markdown
    html = exports.html_export(run)
    assert "<figcaption>Every change needs a reviewer.</figcaption>" in html
    assert 'src="http' in html  # absolute image URL
    document = Document(io.BytesIO(exports.docx_export(run)))
    assert len(document.inline_shapes) == 2  # hero + figure
    assert ("Caption", "Every change needs a reviewer.") in [
        (p.style.name, p.text) for p in document.paragraphs
    ]
    underlined = [r.text for p in document.paragraphs for r in p.runs if r.underline]
    assert "temporary bypasses" in underlined


def test_markdown_export_with_a_brief_fresh_from_the_form(admin_user):
    """A brief that wasn't reloaded keeps its region as a TextChoices member; YAML must still work."""
    brief = ContentBrief.objects.create(topic="MOC", pkeyword="moc", created_by=admin_user)  # default region
    run = services.create_run(brief, {}, user=admin_user)
    for key in ("draft", "cleanup"):
        _approve(run, key, "# MOC\n\nText.")
    assert "region: US" in exports.markdown_export(run)
